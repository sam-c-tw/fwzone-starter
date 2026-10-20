"""讀取 xlsx / csv 為 list[dict]，並寫出輸出 xlsx（SPEC 3.1、第 6 節）。"""

import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ERROR_CODES
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from fwzone.errors import InputError

_READ_ERRORS = (OSError, ValueError, zipfile.BadZipFile, InvalidFileException)
_HEADER_HIGHLIGHT = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")


@dataclass(frozen=True)
class Table:
    """讀入的表格。欄位順序即來源檔順序；所有值皆為字串。"""

    path: Path
    sheet: str | None  # csv 為 None
    columns: list[str]
    rows: list[dict[str, str]] = field(default_factory=list)

    @staticmethod
    def row_number(index: int) -> int:
        """資料列索引（0 起算）對應的來源檔列號（標題列為第 1 列）。"""
        return index + 2

    def location(self, row_number: int) -> str:
        """來源位置，格式 `檔名:工作表!列號`。"""
        if self.sheet is None:
            # TODO(spec): csv 省略工作表時的位置格式未明確定義，目前為 `檔名!列號`
            return f"{self.path.name}!{row_number}"
        return f"{self.path.name}:{self.sheet}!{row_number}"

    def require_columns(self, *names: str) -> None:
        missing = [n for n in names if n not in self.columns]
        if missing:
            where = f"{self.path.name}" + (f" 的工作表 {self.sheet!r}" if self.sheet else "")
            raise InputError(
                f"{where} 找不到欄位 {', '.join(repr(m) for m in missing)}；"
                f"實際欄位：{', '.join(repr(c) for c in self.columns)}"
            )


def _check_readable(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        # TODO(spec): 不支援副檔名的結束碼未定義，目前視為輸入資料錯誤（1）
        raise InputError(f"不支援的檔案格式 {path.suffix!r}（{path}）；僅支援 .xlsx 與 .csv")
    if not path.is_file():
        raise InputError(f"找不到檔案：{path}")
    return suffix


def sheet_names(path: Path) -> list[str]:
    """xlsx 的工作表名稱列表。"""
    _check_readable(path)
    try:
        with pd.ExcelFile(path, engine="openpyxl") as xls:
            return [str(n) for n in xls.sheet_names]
    except _READ_ERRORS as exc:
        raise InputError(f"無法讀取 {path}：{exc}") from exc


def _to_table(path: Path, sheet: str | None, df: pd.DataFrame) -> Table:
    columns = [str(c).strip() for c in df.columns]
    df.columns = pd.Index(columns)
    rows = [{k: str(v) for k, v in rec.items()} for rec in df.to_dict("records")]
    return Table(path, sheet, columns, rows)


def read_table(path: Path, sheet: str | None) -> Table:
    """依副檔名讀取表格。csv 忽略 sheet。所有儲存格以字串讀入，空值為空字串。"""
    suffix = _check_readable(path)

    if suffix == ".csv":
        try:
            df = pd.read_csv(
                path,
                dtype=str,
                keep_default_na=False,
                encoding="utf-8-sig",
                skip_blank_lines=False,
            )
        except UnicodeDecodeError as exc:
            raise InputError(f"{path.name} 無法以 UTF-8 解碼，請另存為 UTF-8 CSV") from exc
        except pd.errors.EmptyDataError:
            return Table(path, None, [])
        except (pd.errors.ParserError, OSError) as exc:
            raise InputError(f"無法讀取 {path}：{exc}") from exc
        return _to_table(path, None, df)

    names = sheet_names(path)
    if sheet not in names:
        raise InputError(
            f"{path.name} 找不到工作表 {sheet!r}；實際工作表：{', '.join(repr(n) for n in names)}"
        )
    try:
        df = pd.read_excel(path, sheet_name=sheet, dtype=str, keep_default_na=False)
    except _READ_ERRORS as exc:
        raise InputError(f"無法讀取 {path}：{exc}") from exc
    return _to_table(path, sheet, df)


@dataclass(frozen=True)
class SheetData:
    """要寫出的工作表。highlight 為要加標題底色的欄位索引（0 起算）。"""

    name: str
    columns: Sequence[str]
    rows: Sequence[Sequence[str]]
    highlight: frozenset[int] = frozenset()


def _text_cell(ws: WriteOnlyWorksheet, value: str) -> WriteOnlyCell | str:
    """以 = 開頭、或剛好是 Excel 錯誤碼（如 #N/A）的字串，要明確標成文字，不被當成公式或錯誤值。"""
    if value.startswith("=") or value in ERROR_CODES:
        cell = WriteOnlyCell(ws, value=value)
        cell.data_type = "s"
        return cell
    return value


def write_xlsx(path: Path, sheets: Sequence[SheetData]) -> None:
    """寫出 xlsx：值一律以字串寫入，標題列凍結並開啟自動篩選。

    使用 openpyxl 的 write-only 模式逐列寫出，資料量大時記憶體用量較低。
    """
    wb = Workbook(write_only=True)
    for sheet in sheets:
        ws = wb.create_sheet(sheet.name)
        ws.freeze_panes = "A2"
        header: list[WriteOnlyCell] = []
        for col_idx, name in enumerate(sheet.columns):
            cell = WriteOnlyCell(ws, value=name)
            cell.data_type = "s"
            if col_idx in sheet.highlight:
                cell.fill = _HEADER_HIGHLIGHT
                cell.font = Font(bold=True)
            header.append(cell)
        ws.append(header)
        for row in sheet.rows:
            ws.append([_text_cell(ws, value) for value in row])
        if sheet.columns:
            last_col = get_column_letter(len(sheet.columns))
            ws.auto_filter.ref = f"A1:{last_col}{len(sheet.rows) + 1}"
    try:
        wb.save(path)
    except PermissionError as exc:
        # TODO(spec): 寫入失敗的結束碼未定義，目前視為輸入資料錯誤（1）
        raise InputError(f"無法寫入 {path}；若檔案正被 Excel 開啟，請先關閉後再執行") from exc

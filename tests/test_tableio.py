import openpyxl
import pytest

from fwzone.errors import InputError
from fwzone.parsing import tokenize
from fwzone.tableio import SheetData, read_table, sheet_names, write_xlsx


def make_xlsx(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


def test_csv_bom_is_stripped(tmp_path):
    path = tmp_path / "a.csv"
    path.write_bytes("﻿source_grp,note\n1.1.1.1,x\n".encode())
    table = read_table(path, None)
    assert table.columns == ["source_grp", "note"]
    assert "﻿" not in table.columns[0]


def test_csv_without_bom(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    assert read_table(path, None).rows == [{"a": "1", "b": "2"}]


def test_csv_non_utf8_hints_to_resave(tmp_path):
    path = tmp_path / "a.csv"
    path.write_bytes("欄位,值\n甲,乙\n".encode("big5"))
    with pytest.raises(InputError, match="UTF-8"):
        read_table(path, None)


def test_values_are_strings_and_empty_is_empty_string(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("id,v\n00123,\nNA,null\n", encoding="utf-8")
    assert read_table(path, None).rows == [
        {"id": "00123", "v": ""},
        {"id": "NA", "v": "null"},
    ]


def test_header_whitespace_is_stripped(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text(" a , b\n1,2\n", encoding="utf-8")
    assert read_table(path, None).columns == ["a", "b"]


def test_multiline_cell_still_tokenizes(tmp_path):
    path = tmp_path / "a.xlsx"
    make_xlsx(path, {"s": [["a"], ["x\r\ny"]]})
    assert tokenize(read_table(path, "s").rows[0]["a"]) == ["x", "y"]


def test_xlsx_reads_requested_sheet_only(tmp_path):
    path = tmp_path / "a.xlsx"
    make_xlsx(path, {"rule": [["a"], ["1"]], "group": [["b"], ["2"]]})
    assert read_table(path, "group").rows == [{"b": "2"}]
    assert sheet_names(path) == ["rule", "group"]


def test_xlsx_missing_sheet_lists_actual_names(tmp_path):
    path = tmp_path / "a.xlsx"
    make_xlsx(path, {"alpha": [["a"]], "beta": [["a"]]})
    with pytest.raises(InputError) as exc:
        read_table(path, "rule")
    assert "'alpha'" in str(exc.value) and "'beta'" in str(exc.value)


def test_csv_ignores_sheet_setting(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("a\n1\n", encoding="utf-8")
    assert read_table(path, "whatever").sheet is None


def test_missing_column_lists_actual_columns(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    table = read_table(path, None)
    with pytest.raises(InputError) as exc:
        table.require_columns("a", "zzz")
    assert "'zzz'" in str(exc.value) and "'a'" in str(exc.value) and "'b'" in str(exc.value)


def test_unsupported_extension_and_missing_file(tmp_path):
    other = tmp_path / "a.txt"
    other.write_text("x", encoding="utf-8")
    with pytest.raises(InputError):
        read_table(other, None)
    with pytest.raises(InputError):
        read_table(tmp_path / "missing.csv", None)


def test_extension_is_case_insensitive(tmp_path):
    path = tmp_path / "A.CSV"
    path.write_text("a\n1\n", encoding="utf-8")
    assert read_table(path, None).rows == [{"a": "1"}]


def test_location_format(tmp_path):
    xlsx = tmp_path / "z.xlsx"
    make_xlsx(xlsx, {"segment": [["a"], ["1"]]})
    assert read_table(xlsx, "segment").location(6) == "z.xlsx:segment!6"


def test_write_xlsx_keeps_error_like_and_formula_like_text(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(path, [SheetData("S", ["a"], [["#N/A"], [""], ["=SUM(1,2)"], ["123"]])])
    assert [r["a"] for r in read_table(path, "S").rows] == ["#N/A", "", "=SUM(1,2)", "123"]
    ws = openpyxl.load_workbook(path)["S"]
    assert [ws.cell(row=i, column=1).data_type for i in (2, 4, 5)] == ["s", "s", "s"]


def test_write_xlsx_header_only_sheet(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(path, [SheetData("S", ["a", "b"], [])])
    table = read_table(path, "S")
    assert table.columns == ["a", "b"] and table.rows == []
    assert openpyxl.load_workbook(path)["S"].auto_filter.ref == "A1:B1"


def test_write_xlsx_round_trip(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(
        path,
        [SheetData("Rules", ["a", "Zone"], [["00123", "=1+1"], ["", "x"]], frozenset({1}))],
    )
    table = read_table(path, "Rules")
    assert table.rows == [{"a": "00123", "Zone": "=1+1"}, {"a": "", "Zone": "x"}]

    ws = openpyxl.load_workbook(path)["Rules"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"
    assert ws["B2"].data_type == "s"
    assert ws["B1"].fill.start_color.rgb != ws["A1"].fill.start_color.rgb
